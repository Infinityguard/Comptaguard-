import io
from decimal import Decimal

import pandas as pd
import pytest

pytest.importorskip("streamlit")
from openpyxl import load_workbook

from ui import excel_bytes, fmt_value, money


def test_money_utilise_des_espaces_insecables():
    assert money(1234567.891, 2) == "1\u00a0234\u00a0567.89"
    assert money(Decimal("1000")) == "1\u00a0000"


def test_fmt_value_types():
    assert fmt_value(12345) == "12\u00a0345"
    assert fmt_value(Decimal("5")) == "5.00"
    assert fmt_value("texte") == "texte"


def test_excel_bytes_met_en_forme_et_neutralise_les_formules():
    df = pd.DataFrame({"nom": ["=CMD()", "Ama"], "net": [Decimal("1500"), Decimal("2500")]})
    wb = load_workbook(io.BytesIO(excel_bytes({"Paie": df})))
    ws = wb["Paie"]
    assert ws.freeze_panes == "A2"
    assert ws["A2"].value == "'=CMD()"          # jamais interprété comme formule
    assert ws["A2"].data_type != "f"
    assert ws["B2"].value == 1500
    assert ws["A1"].font.bold
