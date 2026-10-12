"""Composants d'interface partagés de Comptaguard (présentation uniquement)."""
from __future__ import annotations

import html
import io
import numbers
from decimal import Decimal

import pandas as pd
import streamlit as st

APP_NAME = "Comptaguard"
APP_TAGLINE = "Contrôle comptable SYSCOHADA"
APP_VERSION = "12"

NBSP = "\u00a0"

# --------------------------------------------------------------------------
# Style global
# --------------------------------------------------------------------------
_CSS = """
<style>
#MainMenu, footer {visibility: hidden;}
.block-container {padding-top: 2.2rem; padding-bottom: 3rem; max-width: 1200px;}
h1 {font-size: 1.75rem !important; font-weight: 650 !important; letter-spacing: -0.01em; padding-bottom: .15rem;}
h2, h3 {font-weight: 600 !important; letter-spacing: -0.005em;}
.cg-sub {color: #5B6573; margin: 0 0 1.4rem 0; font-size: 1rem;}
[data-testid="stSidebar"] {border-right: 1px solid #E3E8EF;}
[data-testid="stMetricValue"] {font-size: 1.55rem; font-weight: 600;}
[data-testid="stMetricLabel"] p {color: #5B6573; font-size: .85rem;}
.stButton > button, .stDownloadButton > button, [data-testid="stFormSubmitButton"] > button {
    border-radius: 6px; font-weight: 500;
}
.cg-brand {display: flex; align-items: center; gap: .8rem; margin: 2rem 0 1.2rem 0;}
.cg-mark {width: 2.6rem; height: 2.6rem; border-radius: 8px; background: #1F4E79; color: #fff;
    display: flex; align-items: center; justify-content: center; font-weight: 700; font-size: 1.3rem;}
.cg-name {font-size: 1.3rem; font-weight: 650; line-height: 1.15; color: #1B2430;}
.cg-tag {font-size: .85rem; color: #5B6573;}
</style>
"""


def inject_css() -> None:
    st.markdown(_CSS, unsafe_allow_html=True)


def brand_header() -> None:
    st.markdown(
        f'<div class="cg-brand"><div class="cg-mark">{html.escape(APP_NAME[0])}</div>'
        f'<div><div class="cg-name">{html.escape(APP_NAME)}</div>'
        f'<div class="cg-tag">{html.escape(APP_TAGLINE)}</div></div></div>',
        unsafe_allow_html=True,
    )


def page_header(title: str, description: str | None = None) -> None:
    st.title(title)
    if description:
        st.markdown(f'<p class="cg-sub">{html.escape(description)}</p>', unsafe_allow_html=True)


def limits(*lines: str, title: str = "Limites et précautions") -> None:
    """Regroupe les avertissements dans un volet repliable plutôt que dans le flux de la page."""
    with st.expander(title, expanded=False):
        for line in lines:
            st.markdown(f"- {line}")


def md_escape(text: object) -> str:
    out = str(text)
    for ch in "\\`*_{}[]<>#|":
        out = out.replace(ch, "\\" + ch)
    return out


# --------------------------------------------------------------------------
# Formats
# --------------------------------------------------------------------------
def money(value, decimals: int = 0) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    return f"{number:,.{decimals}f}".replace(",", NBSP)


def fmt_value(value) -> str:
    if isinstance(value, str) or value is None:
        return "—" if value is None else value
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    if isinstance(value, numbers.Integral):
        return f"{int(value):,}".replace(",", NBSP)
    if isinstance(value, (numbers.Real, Decimal)):
        return money(value, 2)
    return str(value)


def kpis(items) -> None:
    """Rangée d'indicateurs : (libellé, valeur[, aide])."""
    columns = st.columns(len(items))
    for col, item in zip(columns, items):
        with col:
            with st.container(border=True):
                st.metric(item[0], fmt_value(item[1]), help=item[2] if len(item) > 2 else None)


# --------------------------------------------------------------------------
# Tableaux
# --------------------------------------------------------------------------
LABELS = {
    "ligne": "Ligne", "probleme": "Anomalie", "gravite": "Gravité", "detail": "Détail",
    "journal": "Journal", "piece": "Pièce", "libelle": "Libellé", "compte": "Compte",
    "date": "Date", "debit": "Débit", "credit": "Crédit", "score_risque": "Score",
    "niveau_risque": "Niveau de risque", "justification_risque": "Justification",
    "statut": "Statut", "motif": "Motif", "ecart_debit": "Écart débit",
    "ecart_credit": "Écart crédit", "ecart_net": "Écart net",
    "debit_grand_livre": "Débit (grand livre)", "credit_grand_livre": "Crédit (grand livre)",
    "debit_balance": "Débit (balance)", "credit_balance": "Crédit (balance)",
    "matricule": "Matricule", "nom": "Nom", "prenom": "Prénom", "poste": "Poste",
    "type_salarie": "Profil", "periode": "Période", "salaire_base": "Salaire de base",
    "heures_supp": "Heures supplémentaires", "commissions": "Commissions", "primes": "Primes",
    "indemnites": "Indemnités", "rappel_regularisation": "Rappel / régularisation",
    "avantages_nature": "Avantages en nature", "brut_cash": "Brut en espèces",
    "base_sociale": "Base sociale", "base_imposable_parametree": "Base imposable",
    "pension_salarie": "Pension (salarié)", "impot_salarial": "Impôt salarial",
    "autres_retenues_taux": "Autres retenues (taux)", "autres_retenues": "Autres retenues",
    "total_retenues": "Total des retenues", "net_a_payer": "Net à payer",
    "pension_employeur": "Pension (employeur)",
    "prestations_familiales_employeur": "Prestations familiales (employeur)",
    "risques_professionnels_employeur": "Risques professionnels (employeur)",
    "cout_total_employeur": "Coût employeur",
    "id": "N°", "campaign_name": "Campagne", "campaign_date": "Date de comptage",
    "warehouse": "Magasin", "counter": "Compteur", "assigned_to": "Affectée à",
    "second_counter": "Second compteur", "status": "Statut", "notes": "Observations",
    "created_by": "Créé par", "created_at": "Créé le", "updated_at": "Mis à jour le",
    "reviewed_by": "Contrôlé par", "reviewed_at": "Contrôlé le", "review_notes": "Motif de la décision",
    "full_name": "Nom complet", "username": "Identifiant", "role": "Rôle", "active": "Actif",
    "event_at": "Date et heure", "action": "Action", "details": "Détails",
}

_RED = "background-color:#FBE4E4;color:#8A1F1F;font-weight:600"
_AMBER = "background-color:#FFF1D6;color:#7A4B00;font-weight:600"
_GREY = "background-color:#E8EEF5;color:#3B4A5C"
_GREEN = "background-color:#E3F3E8;color:#1E6B3A;font-weight:600"
_CELL_STYLES = {
    "Élevé": _RED, "Haute": _RED, "Rejeté": _RED,
    "Modéré": _AMBER, "Moyenne": _AMBER, "À vérifier": _AMBER, "Écarts à justifier": _AMBER,
    "Faible": _GREY, "Basse": _GREY,
    "Concordant": _GREEN, "Approuvé": _GREEN, "Clôturé": _GREEN,
}
_STYLED_COLUMNS = {"niveau_risque", "gravite", "statut", "status"}


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c) for c in out.columns]
    for col in out.columns:
        if out[col].dtype == object and out[col].map(lambda v: isinstance(v, Decimal)).any():
            out[col] = out[col].map(lambda v: float(v) if isinstance(v, Decimal) else v)
    return out


def _cell_style(value) -> str:
    return _CELL_STYLES.get(str(value), "")


def show_table(df: pd.DataFrame | None, height: int | None = None) -> None:
    if df is None or df.empty:
        st.caption("Aucun élément à afficher.")
        return
    df = _clean(df)
    config = {}
    for col in df.columns:
        label = LABELS.get(col, col.replace("_", " ").strip().capitalize())
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            config[col] = st.column_config.DatetimeColumn(label, format="DD/MM/YYYY")
        elif pd.api.types.is_float_dtype(df[col]):
            config[col] = st.column_config.NumberColumn(label, format="%.2f")
        else:
            config[col] = st.column_config.Column(label)
    data = df
    targets = [c for c in df.columns if c in _STYLED_COLUMNS]
    if targets:
        try:
            styler = df.style
            apply = getattr(styler, "map", None) or styler.applymap
            data = apply(_cell_style, subset=targets)
        except Exception:
            data = df
    kwargs = {"hide_index": True, "column_config": config}
    if height:
        kwargs["height"] = height
    try:
        st.dataframe(data, width="stretch", **kwargs)
    except Exception:
        st.dataframe(data, use_container_width=True, **kwargs)


# --------------------------------------------------------------------------
# Fichiers
# --------------------------------------------------------------------------
def read_table(upload, key: str, sheet_label: str = "Feuille à contrôler"):
    """Lit un CSV ou un classeur Excel téléversé ; retourne (tableau brut, nom de feuille)."""
    from controles import read_uploaded_table

    if upload.name.lower().endswith(".csv"):
        return read_uploaded_table(upload, filename=upload.name), "CSV"
    xls = pd.ExcelFile(upload)
    names = xls.sheet_names
    sheet = st.selectbox(sheet_label, names, key=f"{key}_sheet") if len(names) > 1 else names[0]
    return read_uploaded_table(xls, sheet), sheet


def _safe_cell(value):
    if hasattr(value, "tzinfo") and getattr(value, "tzinfo", None) is not None and hasattr(value, "replace"):
        return value.replace(tzinfo=None)
    # Une chaîne commençant par « = » serait interprétée comme une formule par Excel.
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Classeur Excel mis en forme : en-têtes foncés, volet figé, largeurs ajustées."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            sheet_name = name[:31]
            frame = _clean(frame)
            for col in frame.columns:
                series = frame[col]
                if pd.api.types.is_datetime64_any_dtype(series) and getattr(series.dt, "tz", None) is not None:
                    frame[col] = series = series.dt.tz_localize(None)   # Excel refuse les fuseaux horaires
                if not (pd.api.types.is_numeric_dtype(series) or pd.api.types.is_bool_dtype(series)
                        or pd.api.types.is_datetime64_any_dtype(series)):
                    frame[col] = series.map(_safe_cell)
            frame.to_excel(writer, index=False, sheet_name=sheet_name)
            ws = writer.sheets[sheet_name]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E79")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for idx, col in enumerate(frame.columns, start=1):
                longest = max([len(str(col))] + [len(str(v)) for v in frame[col].head(200)])
                ws.column_dimensions[get_column_letter(idx)].width = min(50, longest + 2)
    return buffer.getvalue()
