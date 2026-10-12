"""Tests des contrôles. Lancer : python test_controles.py  (ou : pytest)"""
import io
import os

import pandas as pd
from openpyxl import load_workbook

from controles import (analyze, make_account_reference_template, make_excel_report,
                       parse_amount, prepare_dataframe, read_account_reference, read_ledger, read_uploaded_table, reconcile_ledger_balance)

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLE = os.path.join(HERE, "exemple_ecritures.xlsx")


def _ledger(rows, journal=False):
    cols = ["date", "piece", "libelle", "compte", "debit", "credit"] + (["journal"] if journal else [])
    return prepare_dataframe(pd.DataFrame(rows, columns=cols))


def test_parse_amount():
    assert parse_amount(None) == (0.0, True)
    assert parse_amount(float("nan")) == (0.0, True)
    assert parse_amount("") == (0.0, True)
    assert parse_amount("-") == (0.0, True)
    assert parse_amount(1500) == (1500.0, True)
    assert parse_amount("85 000") == (85000.0, True)
    assert parse_amount("85\u00a0000,50") == (85000.5, True)
    assert parse_amount("85 000,00 FCFA") == (85000.0, True)
    assert parse_amount("15.000") == (15000.0, True)
    assert parse_amount("15,000") == (15000.0, True)
    assert parse_amount("1.250,75") == (1250.75, True)
    assert parse_amount("1,250.75") == (1250.75, True)
    assert parse_amount("(500)") == (-500.0, True)
    assert parse_amount("-500") == (-500.0, True)
    assert parse_amount("abc")[1] is False
    assert parse_amount("1,2,3")[1] is False
    assert parse_amount(True)[1] is False


def test_dates_iso_et_francaises():
    df = _ledger([("2026-03-05", "P1", "a", "601000", 100, None),
                  ("05/03/2026", "P1", "a", "401000", None, 100),
                  (pd.Timestamp("2026-03-05"), "P1", "a", "401000", 0, None),
                  ("31/02/2026", "P1", "a", "401000", 0, None)])
    d = list(df["date"])
    assert d[0] == d[1] == d[2] == pd.Timestamp("2026-03-05")
    assert pd.isna(d[3])


def test_cellule_vide_pas_d_alerte():
    df = _ledger([("2026-03-05", "P1", "a", "601000", 100, None),
                  ("2026-03-05", "P1", "a", "401000", None, 100)])
    assert analyze(df).empty


def test_exemple():
    df = prepare_dataframe(read_ledger(SAMPLE))
    assert df.attrs["lignes_vides"] == 1
    rep = analyze(df)
    got = {}
    for _, r in rep.iterrows():
        key = int(r["ligne"]) if pd.notna(r["ligne"]) else f"{r['journal']}/{r['piece']}"
        got.setdefault(key, set()).add(r["probleme"])
    dup = "Doublon possible ≤ 7 jours"
    expected = {
        10: {dup}, 11: {dup}, 12: {dup}, 13: {dup},
        14: {"Date invalide ou manquante"},
        16: {"Numéro de pièce manquant"}, 17: {"Numéro de pièce manquant"},
        18: {"Montant rond élevé (à vérifier)"}, 19: {"Montant rond élevé (à vérifier)"},
        22: {"Débit et crédit sur la même ligne"},
        23: {"Aucun montant"},
        24: {"Format de numéro de compte invalide"}, 25: {"Format de numéro de compte invalide"},
        26: {"Débit manquant ou non numérique", "Aucun montant"},
        28: {"Montant négatif"}, 29: {"Montant négatif"},
        "VT/0002": {"Pièce non équilibrée"}, "OD/0011": {"Pièce non équilibrée"},
        "AC/0004": {"Pièce non équilibrée"},
    }
    assert got == expected, (got, expected)
    assert (rep["detail"][rep["ligne"] == 10].iloc[0]).count("12") == 1


def test_en_tetes_majuscules_et_zeros_initiaux():
    raw = read_ledger(SAMPLE)
    assert "0001" in set(prepare_dataframe(raw)["piece"].dropna())


def test_journal_evite_le_masquage():
    rows = [("2026-03-05", "001", "a", "601000", 100, None, "AC"),
            ("2026-03-05", "001", "a", "401000", None, 90, "AC"),     # écart -10 en AC
            ("2026-03-06", "001", "b", "401000", 90, None, "BQ"),
            ("2026-03-06", "001", "b", "521000", None, 100, "BQ")]    # écart +... compense
    avec = analyze(_ledger(rows, journal=True))
    sans = analyze(_ledger([r[:6] for r in rows]))
    assert (avec["probleme"] == "Pièce non équilibrée").sum() == 2
    # Sans journal, la date limite le regroupement : les deux journées restent
    # distinctes et les écarts ne peuvent plus se compenser entre elles.
    assert (sans["probleme"] == "Pièce non équilibrée").sum() == 2


def test_libelle_identique_optionnel():
    rows = [("2026-03-05", "P1", "achat A", "601000", 100, None),
            ("2026-03-06", "P1", "achat B", "601000", 100, None)]
    assert (analyze(_ledger(rows))["probleme"].str.startswith("Doublon")).sum() == 2
    assert (analyze(_ledger(rows), require_same_label=True)["probleme"]
            .str.startswith("Doublon")).sum() == 0


def test_comptes_alphanumeriques_et_classe_9():
    rows = [("2026-03-05", "P1", "a", "411CLIENT", 100, None),
            ("2026-03-05", "P1", "a", "901000", None, 100)]
    assert (analyze(_ledger(rows))["probleme"] == "Format de numéro de compte invalide").sum() == 1
    strict = analyze(_ledger(rows), allow_account_suffix=False)
    assert (strict["probleme"] == "Format de numéro de compte invalide").sum() == 2
    avec9 = analyze(_ledger(rows), account_classes="123456789")
    assert (avec9["probleme"] == "Format de numéro de compte invalide").sum() == 0


def test_referentiel():
    rows = [("2026-03-05", "P1", "a", "601000", 100, None),
            ("2026-03-05", "P1", "a", "401000", None, 100)]
    rep = analyze(_ledger(rows), valid_accounts={"601000"})
    assert (rep["probleme"] == "Compte absent du référentiel fourni").sum() == 1


def test_fichier_vide():
    df = _ledger([(None, None, None, None, None, None)])
    assert df.empty and analyze(df).empty
    make_excel_report(analyze(df), df)


def test_export_excel_pas_de_formule():
    rows = [("2026-03-05", "P1", "=1+1", "601000", 100, None),
            ("2026-03-05", "P1", "x", "401000", None, 50)]
    df = _ledger(rows)
    data = make_excel_report(analyze(df), df, {"fichier": "t.xlsx"})
    wb = load_workbook(io.BytesIO(data))
    assert set(wb.sheetnames) >= {"Synthèse", "Anomalies", "Données contrôlées", "Paramètres"}
    for ws in wb.worksheets:
        assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)


if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for t in tests:
        t()
        print("OK ", t.__name__)
    print(f"{len(tests)} tests réussis")


def test_reutilisation_piece_dates_differentes_sans_journal():
    rows = [
        ("2026-01-05", "001", "opération A", "601000", 100, None),
        ("2026-01-05", "001", "opération A", "401000", None, 90),
        ("2026-02-06", "001", "opération B", "601000", 90, None),
        ("2026-02-06", "001", "opération B", "401000", None, 100),
    ]
    report = analyze(_ledger(rows))
    unbalanced = report[report["probleme"] == "Pièce non équilibrée"]
    assert len(unbalanced) == 2
    assert set(pd.to_datetime(unbalanced["date"]).dt.strftime("%Y-%m-%d")) == {
        "2026-01-05", "2026-02-06"
    }


def test_referentiel_importe_normalise_les_comptes():
    buffer = io.BytesIO()
    pd.DataFrame({" Compte ": [" 401000 ", "601000.0"],
                  "libelle": ["Fournisseurs", "Achats"]}).to_excel(buffer, index=False)
    buffer.seek(0)
    accounts, ref = read_account_reference(buffer)
    assert accounts == {"401000", "601000"}
    assert len(ref) == 2


def test_referentiel_refuse_comptes_dupliques():
    buffer = io.BytesIO()
    pd.DataFrame({"compte": ["401000", "401000.0"]}).to_excel(buffer, index=False)
    buffer.seek(0)
    try:
        read_account_reference(buffer)
    except ValueError as exc:
        assert "double" in str(exc).lower()
    else:
        raise AssertionError("Le référentiel avec doublon aurait dû être refusé")


def test_modele_referentiel_excel():
    data = make_account_reference_template()
    wb = load_workbook(io.BytesIO(data), read_only=True)
    assert wb.sheetnames == ["Comptes", "Instructions"]
    assert [c.value for c in wb["Comptes"][1]][0] == "compte"


def test_aliases_export_comptable():
    from controles import map_common_columns
    raw = pd.DataFrame({"Date écriture": ["2026-01-01"], "N° pièce": ["001"],
                        "Libellé": ["Achat"], "N° compte": ["601000"],
                        "Débit": [100], "Crédit": [0]})
    mapped = map_common_columns(raw)
    assert {"date", "piece", "libelle", "compte", "debit", "credit"}.issubset(mapped.columns)


def test_balance_equilibree_et_desequilibree():
    from controles import analyze_trial_balance
    balanced = pd.DataFrame({"Compte": ["401000", "601000"], "Débit": [100, 0], "Crédit": [0, 100]})
    issues, summary = analyze_trial_balance(balanced)
    assert abs(summary["ecart"]) < 0.01
    assert not (issues["probleme"] == "Balance déséquilibrée").any()
    unbalanced = pd.DataFrame({"Compte": ["401000", "601000"], "Solde débiteur": [100, 0], "Solde créditeur": [0, 90]})
    issues, summary = analyze_trial_balance(unbalanced)
    assert abs(summary["ecart"] - 10) < 0.01
    assert (issues["probleme"] == "Balance déséquilibrée").any()


def test_balance_comptes_dupliques_signales():
    from controles import analyze_trial_balance
    raw = pd.DataFrame({"compte": ["401000", "401000"], "debit": [50, 50], "credit": [0, 0]})
    issues, _ = analyze_trial_balance(raw)
    assert (issues["probleme"] == "Compte présent plusieurs fois").any()


def test_rate_risks_prioritizes_and_explains_alerts():
    from controles import rate_risks
    import pandas as pd

    alerts = pd.DataFrame([
        {"probleme": "Pièce non équilibrée", "gravite": "Haute", "debit": 12_000_000, "credit": 0},
        {"probleme": "Montant rond élevé (à vérifier)", "gravite": "Basse", "debit": 100_000, "credit": 0},
    ])
    rated = rate_risks(alerts)
    assert rated.loc[0, "niveau_risque"] == "Élevé"
    assert rated.loc[0, "score_risque"] > rated.loc[1, "score_risque"]
    assert rated["justification_risque"].str.len().gt(0).all()


def test_rate_risks_handles_empty_report():
    from controles import rate_risks
    import pandas as pd

    rated = rate_risks(pd.DataFrame(columns=["probleme", "gravite"]))
    assert rated.empty
    assert "score_risque" in rated.columns
    assert "niveau_risque" in rated.columns


def test_reconcile_ledger_balance_matching_and_mismatch():
    ledger = pd.DataFrame({
        "compte": ["401", "411", "401"],
        "debit": [0, 100, 50],
        "credit": [50, 0, 0],
    })
    balance = pd.DataFrame({
        "compte": ["401", "411"],
        "debit": [60, 100],
        "credit": [50, 0],
    })
    result, summary = reconcile_ledger_balance(ledger, balance)
    assert summary["comptes_comparés"] == 2
    assert summary["comptes_concordants"] == 1
    assert summary["comptes_a_verifier"] == 1
    assert result.loc[result["compte"] == "401", "statut"].iloc[0] == "À vérifier"
    assert result.loc[result["compte"] == "411", "statut"].iloc[0] == "Concordant"


def test_reconcile_rejects_balance_saldos_only():
    ledger = pd.DataFrame({"compte": ["401"], "debit": [0], "credit": [100]})
    balance = pd.DataFrame({"compte": ["401"], "solde debiteur": [0], "solde crediteur": [100]})
    try:
        reconcile_ledger_balance(ledger, balance)
    except ValueError as exc:
        assert "mouvements débit et crédit" in str(exc)
    else:
        raise AssertionError("Une balance en soldes ne doit pas être comparée directement aux mouvements.")


def test_import_csv_point_virgule_et_zeros_initiaux():
    csv_data = ("Date écriture;N° pièce;Libellé;N° compte;Débit;Crédit\n"
                "2026-01-02;0007;Achat café;060100;15 000,50;0\n"
                "2026-01-02;0007;Fournisseur;040100;0;15 000,50\n").encode("utf-8-sig")
    buffer = io.BytesIO(csv_data)
    buffer.name = "ecritures.csv"
    raw = read_uploaded_table(buffer)
    assert {"date", "piece", "libelle", "compte", "debit", "credit"}.issubset(raw.columns)
    assert raw.loc[0, "piece"] == "0007"
    assert raw.loc[0, "compte"] == "060100"
    assert raw.loc[0, "debit"] == "15 000,50"


def test_import_csv_virgule_et_colonnes_communes():
    csv_data = b"date,piece,libelle,compte,debit,credit\n2026-01-02,P1,Achat,601000,100,0\n"
    buffer = io.BytesIO(csv_data)
    buffer.name = "journal.csv"
    raw = read_uploaded_table(buffer)
    assert raw.iloc[0]["compte"] == "601000"
    assert raw.iloc[0]["debit"] == "100"


def test_import_csv_rejette_entetes_non_reconnues_en_controle():
    csv_data = b"jour;reference;description;numero;sortie;entree\n2026-01-02;P1;Achat;601000;100;0\n"
    buffer = io.BytesIO(csv_data)
    buffer.name = "inconnu.csv"
    raw = read_uploaded_table(buffer)
    try:
        prepare_dataframe(raw)
    except ValueError as exc:
        assert "manquantes" in str(exc).lower()
    else:
        raise AssertionError("Des colonnes obligatoires non reconnues auraient dû être signalées")


def test_qualite_import_detecte_colonnes_et_montants():
    from controles import assess_import_quality
    raw = pd.DataFrame({"Date": ["date invalide"], "Pièce": ["P1"], "Libellé": ["Achat"],
                        "Compte": [""], "Débit": ["abc"], "Crédit": [0]})
    report, summary = assess_import_quality(raw, "Écritures comptables")
    assert summary["nombre_lignes"] == 1
    assert summary["nombre_constats"] >= 3
    assert "Montants illisibles : debit" in set(report["constat"])
    assert "Dates à vérifier" in set(report["constat"])


def test_qualite_import_detecte_colonnes_manquantes():
    from controles import assess_import_quality
    report, summary = assess_import_quality(pd.DataFrame({"Compte": ["401000"], "Débit": [100]}), "Balance générale")
    assert summary["colonnes_manquantes"] == ["credit"]
    assert "Colonnes nécessaires absentes" in set(report["constat"])
